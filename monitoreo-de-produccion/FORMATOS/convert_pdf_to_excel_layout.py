import pdfplumber
import pandas as pd
import re
from openpyxl import Workbook

pdf_path = "RegistroCondicionesProduccion_OsmosisInversa.pdf"
excel_path = "RegistroCondicionesProduccion_OsmosisInversa_Completo.xlsx"

wb = Workbook()
# Remove default sheet
wb.remove(wb.active)

with pdfplumber.open(pdf_path) as pdf:
    for i, page in enumerate(pdf.pages):
        ws = wb.create_sheet(title=f"Pagina {i+1}")
        
        # Extract text preserving visual layout
        text = page.extract_text(layout=True)
        if not text:
            continue
            
        lines = text.split('\n')
        
        row_idx = 1
        for line in lines:
            line_stripped = line.strip()
            if not line_stripped:
                # Optionally skip empty lines or keep them for spacing
                row_idx += 1
                continue
                
            # Split by 2 or more spaces
            parts = re.split(r'\s{2,}', line_stripped)
            
            for col_idx, part in enumerate(parts):
                ws.cell(row=row_idx, column=col_idx+1, value=part.strip())
                
            row_idx += 1

wb.save(excel_path)
print(f"Successfully converted PDF preserving layout to {excel_path}")
