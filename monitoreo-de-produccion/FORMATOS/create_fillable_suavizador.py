import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

def create_fillable_suavizador(filename="Bitacora_Suavizador_Digital.xlsx"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bitacora Suavizador"

    # Styles
    bold_font = Font(bold=True, size=11)
    header_font = Font(bold=True, size=14)
    normal_font = Font(size=10)
    
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
    right_align = Alignment(horizontal="right", vertical="center", wrap_text=True)
    
    thin_border = Border(left=Side(style='thin'), 
                         right=Side(style='thin'), 
                         top=Side(style='thin'), 
                         bottom=Side(style='thin'))
                         
    fill_bg = PatternFill(start_color="E3F2FD", end_color="E3F2FD", fill_type="solid") # Light blue for input
    title_bg = PatternFill(start_color="CFD8DC", end_color="CFD8DC", fill_type="solid") # Light blue-gray for headers
    
    # Configure Column Widths
    ws.column_dimensions['A'].width = 15 # Fecha
    ws.column_dimensions['B'].width = 15 # Suavizador
    ws.column_dimensions['C'].width = 10 # Inicio
    ws.column_dimensions['D'].width = 10 # Final
    ws.column_dimensions['E'].width = 10 # T. Total
    # Hours columns
    for col in 'FGHIJKLMNOPQRSTU':
        ws.column_dimensions[col].width = 8

    # Title
    ws.merge_cells('A1:U1')
    ws['A1'] = "4 MONITOREO DE TRABAJO DEL SUAVIZADOR"
    ws['A1'].font = header_font
    ws['A1'].alignment = center_align
    ws['A1'].fill = title_bg

    # Header Info
    ws['A3'] = "Planta:"
    ws['A3'].font = bold_font
    ws['B3'].fill = fill_bg
    ws['B3'].border = thin_border
    
    ws['D3'] = "Mes:"
    ws['D3'].font = bold_font
    ws['E3'].fill = fill_bg
    ws['E3'].border = thin_border
    
    ws['G3'] = "Año:"
    ws['G3'].font = bold_font
    ws['H3'].fill = fill_bg
    ws['H3'].border = thin_border

    # Table Headers - Row 5 & 6
    headers_main = ["Fecha", "Num Suav.", "Inicio", "Final", "T. Total"]
    hours = [f"{h}:00" for h in range(6, 22)]
    
    for i, h in enumerate(headers_main):
        cell = ws.cell(row=5, column=i+1, value=h)
        cell.font = bold_font
        cell.alignment = center_align
        cell.border = thin_border
        cell.fill = title_bg
        ws.merge_cells(start_row=5, start_column=i+1, end_row=6, end_column=i+1)

    ws.merge_cells('F5:U5')
    ws['F5'] = "MONITOREO DE TRABAJO DEL SUAVIZADOR (5 - 30 GOTAS)"
    ws['F5'].font = bold_font
    ws['F5'].alignment = center_align
    ws['F5'].border = thin_border
    ws['F5'].fill = title_bg

    for i, h in enumerate(hours):
        cell = ws.cell(row=6, column=i+6, value=h)
        cell.font = bold_font
        cell.alignment = center_align
        cell.border = thin_border
        cell.fill = title_bg

    # Data Rows for Monitoreo (e.g., 31 days)
    row = 7
    for _ in range(31):
        for col_idx in range(1, 22):
            cell = ws.cell(row=row, column=col_idx)
            cell.border = thin_border
            cell.fill = fill_bg
        row += 1

    row += 2
    # Section 2: Ciclos de Lavado
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=15)
    ws.cell(row=row, column=1, value="BITÁCORA DE REGENERACIÓN Y LAVADO").font = header_font
    ws.cell(row=row, column=1).alignment = center_align
    ws.cell(row=row, column=1).fill = title_bg
    row += 1

    sub_headers = ["Num Suav.", "Retro lavado", "Regenerado", "Enjuage Lento", "Enjuage Rápido", "T. Total", "Firma"]
    cols_span = [1, 3, 3, 3, 3, 1, 1]
    
    current_col = 1
    for h, span in zip(sub_headers, cols_span):
        cell = ws.cell(row=row, column=current_col, value=h)
        cell.font = bold_font
        cell.alignment = center_align
        cell.border = thin_border
        cell.fill = title_bg
        if span > 1:
            ws.merge_cells(start_row=row, start_column=current_col, end_row=row, end_column=current_col + span - 1)
        current_col += span
    
    row += 1
    sub_sub_headers = ["Tiempo", "Inicio", "Final"]
    current_col = 2
    for _ in range(4): # For the 4 cycles
        for h in sub_sub_headers:
            cell = ws.cell(row=row, column=current_col, value=h)
            cell.font = bold_font
            cell.alignment = center_align
            cell.border = thin_border
            cell.fill = title_bg
            current_col += 1

    row += 1
    # Data Rows for Cycles
    for _ in range(10):
        for col_idx in range(1, 16):
            cell = ws.cell(row=row, column=col_idx)
            cell.border = thin_border
            cell.fill = fill_bg
        row += 1

    # Protection
    for r in ws.iter_rows():
        for cell in r:
            if cell.fill == fill_bg:
                cell.protection = openpyxl.styles.Protection(locked=False)
    
    ws.protection.enable()

    wb.save(filename)
    print(f"Fillable form created successfully: {filename}")

if __name__ == "__main__":
    create_fillable_suavizador()
