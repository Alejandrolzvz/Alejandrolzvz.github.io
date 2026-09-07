import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

def create_fillable_calidad(filename="Bitacora_Control_Calidad_Digital.xlsx"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Control de Calidad"

    # Styles
    bold_font = Font(bold=True, size=11)
    header_font = Font(bold=True, size=14)
    normal_font = Font(size=10)
    
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
    right_align = Alignment(horizontal="right", vertical="center", wrap_text=True)
    
    thin_border = Border(left=Side(style='thin'), 
                         right=Side(style='thin'), 
                         top=Side(style='thin'), 
                         bottom=Side(style='thin'))
                         
    fill_bg = PatternFill(start_color="E1F5FE", end_color="E1F5FE", fill_type="solid") # Light blue for inputs
    title_bg = PatternFill(start_color="CFD8DC", end_color="CFD8DC", fill_type="solid") # Gray-blue for headers
    
    # Configure Column Widths
    ws.column_dimensions['A'].width = 35 # Parameter Name
    # Hours columns (6:00 to 21:00 -> 16 columns)
    hours = [f"{h}:00" for h in range(6, 22)]
    for i, _ in enumerate(hours):
        col_letter = openpyxl.utils.get_column_letter(i + 2)
        ws.column_dimensions[col_letter].width = 10

    # Main Title
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(hours)+1)
    ws['A1'] = "BITÁCORA DE CONTROL DE CALIDAD (MONITOREO)"
    ws['A1'].font = header_font
    ws['A1'].alignment = center_align
    ws['A1'].fill = title_bg

    # Header Info - Row 3
    headers_info = [("Planta:", "B3"), ("Fecha:", "E3"), ("Responsable 1:", "H3"), ("Responsable 2:", "K3"), ("Supervisor:", "O3")]
    for label, cell_ref in headers_info:
        cell = ws[cell_ref]
        cell.value = label
        cell.font = bold_font
        # Set input neighbor
        input_cell = ws.cell(row=ws[cell_ref].row, column=ws[cell_ref].column + 1)
        input_cell.fill = fill_bg
        input_cell.border = thin_border
        if label == "Responsable 1:": # Span names a bit
             ws.merge_cells(start_row=3, start_column=9, end_row=3, end_column=10)
             ws.cell(row=3, column=9).fill = fill_bg
             ws.cell(row=3, column=9).border = thin_border
        elif label == "Responsable 2:":
             ws.merge_cells(start_row=3, start_column=12, end_row=3, end_column=13)
             ws.cell(row=3, column=12).fill = fill_bg
             ws.cell(row=3, column=12).border = thin_border

    # Table Header - Row 5
    ws['A5'] = "PARÁMETROS / HORA"
    ws['A5'].font = bold_font
    ws['A5'].alignment = center_align
    ws['A5'].fill = title_bg
    ws['A5'].border = thin_border

    for i, h in enumerate(hours):
        cell = ws.cell(row=5, column=i+2, value=h)
        cell.font = bold_font
        cell.alignment = center_align
        cell.fill = title_bg
        cell.border = thin_border

    # Section Helper
    def add_section(start_row, title, parameters):
        ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=len(hours)+1)
        ws.cell(row=start_row, column=1, value=title).font = bold_font
        ws.cell(row=start_row, column=1).fill = title_bg
        ws.cell(row=start_row, column=1).border = thin_border
        ws.cell(row=start_row, column=1).alignment = left_align
        
        current_row = start_row + 1
        for param in parameters:
            c1 = ws.cell(row=current_row, column=1, value=param)
            c1.font = normal_font
            c1.border = thin_border
            for col_idx in range(2, len(hours) + 2):
                c = ws.cell(row=current_row, column=col_idx)
                c.fill = fill_bg
                c.border = thin_border
            current_row += 1
        return current_row

    row = 6
    row = add_section(row, "Cisterna Cruda (Proceso A)", [
        "Cloro Residual (2 ppm)",
        "pH (6.8 - 7.5)"
    ])
    
    row += 1
    row = add_section(row, "Cisterna Pretratada (Proceso B)", [
        "Cloro Residual (0.2 ppm)",
        "pH (6.8 - 7.5)"
    ])

    row += 1
    row = add_section(row, "Tanque de Almacenamiento / Muestra Final", [
        "pH (6.8 - 7.5 Unidades)",
        "S.T.D. (80 - 90 ppm)",
        "Conductividad (160 - 180 µ/S cm)",
        "Dureza (5 - 10 ppm)",
        "Ozono (3 - 4 gotas)",
        "Cloro Residual (0 ppm)"
    ])

    row += 1
    row = add_section(row, "Lavado de Envases (Lavado Alcalino)", [
        "pH (10 - 10.5 Unidades)",
        "Concentración G3 (0.4 - 0.5%)"
    ])

    row += 2
    # Supplemental Section: Lamps and Hours
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=17)
    ws.cell(row=row, column=1, value="EQUIPOS Y LÁMPARAS GERMICIDAS").font = bold_font
    ws.cell(row=row, column=1).fill = title_bg
    ws.cell(row=row, column=1).border = thin_border
    row += 1
    
    equipos = [
        "Lámpara Germicida #1 (Hora Apagado / Hora Encendido)",
        "Lámpara Germicida #2 (Hora Apagado / Hora Encendido)",
        "Récord de Horas Equipo 1",
        "Récord de Horas Equipo 2"
    ]
    for eq in equipos:
        ws.cell(row=row, column=1, value=eq).border = thin_border
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)
        ws.cell(row=row, column=2).fill = fill_bg
        ws.cell(row=row, column=2).border = thin_border
        row += 1

    # Protection
    for r in ws.iter_rows():
        for cell in r:
            if cell.fill == fill_bg:
                cell.protection = openpyxl.styles.Protection(locked=False)
                
    ws.protection.enable()

    wb.save(filename)
    print(f"Fillable quality form created successfully: {filename}")

if __name__ == "__main__":
    create_fillable_calidad()
