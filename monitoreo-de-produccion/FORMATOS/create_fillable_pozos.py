import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

def create_fillable_pozos(filename="Registro_Calidad_Pozos_Digital.xlsx"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Registro Calidad Pozos"

    # Styles
    bold_font = Font(bold=True, size=11)
    header_font = Font(bold=True, size=14)
    normal_font = Font(size=10)
    
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
    right_align = Alignment(horizontal="right", vertical="center", wrap_text=True)
    
    thin_border = Border(left=Side(style='thin'), 
                         right=Side(style='thin'), 
                         top=Side(style='thin'), 
                         bottom=Side(style='thin'))
                         
    fill_bg = PatternFill(start_color="E1F5FE", end_color="E1F5FE", fill_type="solid") # Light blue for inputs
    title_bg = PatternFill(start_color="CFD8DC", end_color="CFD8DC", fill_type="solid") # Gray-blue for headers
    
    # Configure Column Widths
    cols_widths = {
        'A': 8,   # Pozo
        'B': 8,   # Día
        'C': 15,  # Conductividad
        'D': 15,  # S.T.D.
        'E': 15,  # pH
        'F': 12,  # Cloro 7:00
        'G': 12,  # Cloro 1:00
        'H': 12,  # Cloro 5:00
        'I': 30,  # Firma Jefe Turno
        'J': 30   # Firma Producción
    }
    for col, width in cols_widths.items():
        ws.column_dimensions[col].width = width

    # Main Title
    ws.merge_cells('A1:J1')
    ws['A1'] = "1 REGISTRO DE CALIDAD DE POZOS"
    ws['A1'].font = header_font
    ws['A1'].alignment = center_align
    ws['A1'].fill = title_bg

    # Header Info - Row 3
    ws['A3'] = "Planta:"
    ws['A3'].font = bold_font
    ws['B3'].fill = fill_bg
    ws['B3'].border = thin_border
    
    ws['D3'] = "Mes:"
    ws['D3'].font = bold_font
    ws['E3'].fill = fill_bg
    ws['E3'].border = thin_border
    
    ws['G3'] = "Año:"
    ws['G3'].font = bold_font
    ws['H3'].fill = fill_bg
    ws['H3'].border = thin_border

    # Table Header - Row 5
    headers = [
        ("Pozo", 'A5'), ("Día", 'B5'), ("Conductividad", 'C5'), ("S.T.D.", 'D5'), 
        ("pH", 'E5'), ("Cloro 7:00 AM", 'F5'), ("Cloro 1:00 PM", 'G5'), 
        ("Cloro 5:00 PM", 'H5'), ("Firma Jefe Turno", 'I5'), ("Firma Prod/Calidad", 'J5')
    ]
    
    for text, cell_ref in headers:
        cell = ws[cell_ref]
        cell.value = text
        cell.font = bold_font
        cell.alignment = center_align
        cell.fill = title_bg
        cell.border = thin_border

    # Sub-headers (Spec ranges) - Row 6
    specs = [
        ("", "A6"), ("", "B6"), ("(500-1500)", "C6"), ("(700-1000) ppm", "D6"),
        ("(6.5 - 8.5) Unid", "E6"), ("(2 ppm)", "F6"), ("(2 ppm)", "G6"),
        ("(2 ppm)", "H6"), ("", "I6"), ("", "J6")
    ]
    for text, cell_ref in specs:
        cell = ws[cell_ref]
        cell.value = text
        cell.font = normal_font
        cell.alignment = center_align
        cell.fill = title_bg
        cell.border = thin_border

    # Data Rows
    # Replicating the PDF structure: Multiple pools/pozos, each with potentially multiple days or just space
    row = 7
    pozos = ["1", "2", "3", "4", "5", "6", "7", "8"] # Example number of wells
    for p in pozos:
        # Each well could have 3 rows for more space or just 1. Let's do 1 for clean layout as Excel is scrollable.
        for col_idx in range(1, 11):
            cell = ws.cell(row=row, column=col_idx)
            cell.border = thin_border
            cell.fill = fill_bg
            if col_idx == 1:
                cell.value = p
                cell.font = bold_font
                cell.alignment = center_align
                cell.fill = title_bg # Wells column stays gray
        row += 1

    # More space for more entries
    for _ in range(20):
        for col_idx in range(1, 11):
            cell = ws.cell(row=row, column=col_idx)
            cell.border = thin_border
            cell.fill = fill_bg
        row += 1

    # Protection
    for r in ws.iter_rows():
        for cell in r:
            if cell.fill == fill_bg:
                cell.protection = openpyxl.styles.Protection(locked=False)
                
    ws.protection.enable()

    wb.save(filename)
    print(f"Fillable wells quality form created successfully: {filename}")

if __name__ == "__main__":
    create_fillable_pozos()
